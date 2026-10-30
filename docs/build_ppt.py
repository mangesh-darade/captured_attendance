from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

TEAL = RGBColor(0x0E, 0x6B, 0x5C)
DARK = RGBColor(0x1F, 0x3A, 0x33)
CREAM = RGBColor(0xF7, 0xF6, 0xF2)
MINT = RGBColor(0xE7, 0xF2, 0xEE)
WHITE = RGBColor(255, 255, 255)
MUTED = RGBColor(0x5F, 0x6F, 0x6A)
AMBER = RGBColor(0xB4, 0x6B, 0x00)
BLUE = RGBColor(0x2F, 0x7C, 0xF6)
PURPLE = RGBColor(0x7C, 0x3A, 0xED)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
blank = prs.slide_layouts[6]


def bg(slide, color):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = color


def box(slide, x, y, w, h, fill=WHITE, line=None, radius=True):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line:
        s.line.color.rgb = line
        s.line.width = Pt(1.25)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    if radius:
        s.adjustments[0] = 0.08
    return s


def text(slide, x, y, w, h, t, size=16, bold=False, color=DARK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    lines = t if isinstance(t, list) else [t]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = 'Calibri'
    return tb


def title(slide, t, sub=None):
    bg(slide, CREAM)
    bar = box(slide, 0, 0, 13.333, 0.18, TEAL, radius=False)
    text(slide, 0.6, 0.4, 12, 0.8, t, 32, True, TEAL)
    if sub:
        text(slide, 0.6, 1.15, 12, 0.5, sub, 16, False, MUTED)


def card(slide, x, y, w, h, head, body, accent=TEAL, head_size=17, body_size=13):
    box(slide, x, y, w, h, WHITE, MINT)
    box(slide, x, y, 0.09, h, accent, radius=False)
    text(slide, x + 0.25, y + 0.12, w - 0.4, 0.5, head, head_size, True, accent)
    text(slide, x + 0.25, y + 0.62, w - 0.4, h - 0.7, body, body_size, False, DARK)


def footer(slide, n):
    text(slide, 0.6, 7.05, 8, 0.3, 'ElintOmAT  |  Attendance & Work Companion', 10, False, MUTED)
    text(slide, 12.0, 7.05, 0.8, 0.3, str(n), 10, False, MUTED, PP_ALIGN.RIGHT)


n = 0


def new(t, sub=None):
    global n
    n += 1
    s = prs.slides.add_slide(blank)
    title(s, t, sub)
    footer(s, n)
    return s


# 1 title
n += 1
s = prs.slides.add_slide(blank)
bg(s, TEAL)
text(s, 0.9, 2.2, 11.5, 1.2, 'ElintOmAT', 60, True, WHITE)
text(s, 0.9, 3.5, 11.5, 0.8, 'Your personal work companion for attendance, visits and approvals', 24, False, RGBColor(0xCF, 0xE8, 0xE0))
text(s, 0.9, 5.6, 11.5, 0.5, 'Android app  +  Web for HR / Managers  |  Product overview and user guide', 16, False, RGBColor(0xCF, 0xE8, 0xE0))

# 2 why
s = new('Why ElintOmAT?', 'Attendance should take 5 seconds, not a conversation.')
card(s, 0.6, 1.9, 3.9, 2.4, 'Employees', ['Check in with a face scan', 'Tell the manager "I will be late" in 2 taps', 'See my hours, leave balance and requests'])
card(s, 4.7, 1.9, 3.9, 2.4, 'Managers / HR', ['Approve WFH and leave on the web', 'See who is where, and why', 'Plan field visits and send announcements'], BLUE)
card(s, 8.8, 1.9, 3.9, 2.4, 'Business', ['Trusted attendance (face + GPS)', 'Less follow-up and paperwork', 'Clear records for payroll and audits'], PURPLE)
box(s, 0.6, 4.7, 12.1, 1.6, MINT)
text(s, 0.9, 4.85, 11.5, 1.4, ['One system: the app for employees, the web for managers, one backend for everything.',
                              'Works in English, Marathi, Hindi and Gujarati.'], 18, False, DARK)

# 3 overview
s = new('What is in the product', 'Three parts that work together')
card(s, 0.6, 1.9, 3.9, 3.3, 'Android app', ['Home with one-tap check-in', 'Quick updates to the manager', 'Field-visit tools', 'Records and My day', 'Reminders and location'])
card(s, 4.7, 1.9, 3.9, 3.3, 'Web (HR module)', ['Attendance requests inbox', 'Announcements', 'App Home messages', 'Visit plans', 'Attendance reports'], BLUE)
card(s, 8.8, 1.9, 3.9, 3.3, 'Backend + API', ['Secure token login', 'Face matching on the server', 'Site / geofence rules', 'Leave balances and shifts', 'One database for all'], PURPLE)
text(s, 0.6, 5.6, 12, 0.6, 'App  <->  secure API  <->  database  <->  web screens', 18, True, TEAL, PP_ALIGN.CENTER)

# 4 home
s = new('The Home screen', 'Everything important in one view, no endless scrolling')
box(s, 0.6, 1.8, 3.6, 5.0, WHITE, MINT)
text(s, 0.8, 1.95, 3.2, 0.4, 'Good morning, Rahul', 15, True, TEAL)
for i, (lbl, c) in enumerate([('Running late', AMBER), ('Work from home', PURPLE), ("Can't come", RGBColor(0xB9, 0x1C, 0x1C)), ('Take leave', BLUE)]):
    b = box(s, 0.8 + i * 0.82, 2.55, 0.76, 0.55, WHITE, c)
    text(s, 0.8 + i * 0.82, 2.58, 0.76, 0.5, lbl, 8, True, c, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
box(s, 0.8, 3.25, 3.2, 1.7, MINT)
ci = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.95), Inches(3.4), Inches(1.3), Inches(1.3))
ci.fill.solid(); ci.fill.fore_color.rgb = TEAL; ci.line.fill.background()
text(s, 0.95, 3.75, 1.3, 0.6, 'Check in', 14, True, WHITE, PP_ALIGN.CENTER)
text(s, 2.35, 3.5, 1.6, 1.0, ['Pyramid County', 'View on map'], 10, False, DARK)
box(s, 0.8, 5.1, 3.2, 0.6, RGBColor(0xFF, 0xF3, 0xD6))
text(s, 0.9, 5.17, 3.0, 0.5, 'Announcement: Happy Diwali', 11, True, AMBER)
box(s, 0.8, 5.85, 3.2, 0.7, MINT)
text(s, 0.9, 5.9, 3.0, 0.6, '"Consistency beats intensity."', 11, False, TEAL)
card(s, 4.6, 1.8, 4.0, 1.5, 'Check-in card', 'Status, shift time, big Check In button and live location with a map shortcut.', TEAL, 15, 12)
card(s, 4.6, 3.45, 4.0, 1.5, 'Quick tiles', 'Late, WFH, can\'t come, leave: the manager is told in two taps.', AMBER, 15, 12)
card(s, 4.6, 5.1, 4.0, 1.5, 'Stats', 'Days this week, on-time streak and month attendance %.', BLUE, 15, 12)
card(s, 8.8, 1.8, 3.9, 1.5, 'Announcements', 'Swipe through notices; acknowledge important ones.', PURPLE, 15, 12)
card(s, 8.8, 3.45, 3.9, 1.5, "Today's thought", 'A line of motivation chosen by HR on the web.', TEAL, 15, 12)
card(s, 8.8, 5.1, 3.9, 1.5, 'Daily heading', 'The big line at the top is also managed from the web.', BLUE, 15, 12)

# 5 check in
s = new('Check-in in seconds', 'Face + location, with a safe fallback for every situation')
steps = [('1', 'Tap Check In', 'Camera opens with a face guide'), ('2', 'Face scan', 'Matched on the server (128-point face data, no photo stored for matching)'),
         ('3', 'Location', 'Live GPS matched to saved sites. Unknown site is saved, not blocked'), ('4', 'Done', 'Vibration + success message; Home shows the visit timer')]
for i, (num, h, b) in enumerate(steps):
    x = 0.6 + i * 3.1
    box(s, x, 1.9, 2.9, 2.5, WHITE, MINT)
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.2), Inches(2.05), Inches(0.6), Inches(0.6))
    c.fill.solid(); c.fill.fore_color.rgb = TEAL; c.line.fill.background()
    text(s, x + 0.2, 2.12, 0.6, 0.5, num, 20, True, WHITE, PP_ALIGN.CENTER)
    text(s, x + 0.2, 2.8, 2.6, 0.4, h, 17, True, TEAL)
    text(s, x + 0.2, 3.2, 2.6, 1.2, b, 12, False, DARK)
card(s, 0.6, 4.7, 4.0, 2.0, 'Face fails 3 times?', 'The app offers "Check in without face scan". The punch is saved and the manager gets a note to review.', AMBER, 16, 12)
card(s, 4.7, 4.7, 4.0, 2.0, 'No internet?', 'The punch is saved as Pending with the real time, and syncs automatically. You see "Synced" when done.', BLUE, 16, 12)
card(s, 8.8, 4.7, 3.9, 2.0, 'Working from home?', 'If the manager approved WFH for today, the punch is marked WFH automatically.', PURPLE, 16, 12)

# 6 quick updates
s = new('Tell the manager in two taps', 'Late, WFH, can\'t come and leave, with approval and Undo')
flow = [('Employee', 'Taps a tile, picks time / reason / dates, adds note or a voice note'), ('Manager', 'Gets it in the web inbox and a notification'),
        ('Decision', 'Approve or reject with a note'), ('Employee', 'Sees Approved / Rejected on Home; leave balance updates')]
for i, (h, b) in enumerate(flow):
    x = 0.6 + i * 3.1
    box(s, x, 2.0, 2.8, 2.3, WHITE, TEAL)
    text(s, x + 0.2, 2.15, 2.4, 0.4, h, 18, True, TEAL)
    text(s, x + 0.2, 2.7, 2.4, 1.6, b, 13, False, DARK)
    if i < 3:
        a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + 2.82), Inches(2.95), Inches(0.26), Inches(0.4))
        a.fill.solid(); a.fill.fore_color.rgb = TEAL; a.line.fill.background()
card(s, 0.6, 4.7, 3.9, 1.9, 'Status stays visible', 'Sent, Pending, Approved, Rejected. Undo any time before approval.', TEAL, 16, 12)
card(s, 4.7, 4.7, 3.9, 1.9, 'Smart clearing', '"Running late" clears itself when you check in.', BLUE, 16, 12)
card(s, 8.8, 4.7, 3.9, 1.9, 'Missed a check-out?', 'A card on Home lets you request a correction in 10 seconds.', AMBER, 16, 12)

# 7 field visits
s = new('Field visits made simple', 'Plan from the web, work from the phone')
card(s, 0.6, 1.9, 4.0, 2.3, "Today's plan", ['Admin plans the stops for each employee on the web', 'The app lists them: Done, Now, next', 'End visit moves to the next stop'], TEAL, 17, 13)
card(s, 4.7, 1.9, 4.0, 2.3, 'Visit card', ['Customer name and address', 'Open orders, last visit date, tasks due', 'Live timer from check-in'], BLUE, 17, 13)
card(s, 8.8, 1.9, 3.9, 2.3, 'Capture on the spot', ['Add note', 'Photo', 'Voice note (up to 60 s)', 'Details with all of today\'s notes'], PURPLE, 17, 13)
box(s, 0.6, 4.6, 12.1, 2.0, MINT)
text(s, 0.9, 4.75, 11.5, 1.8, ['After the last visit: Day complete shows worked time, number of visits and time spent outside the site area.',
                              'The employee adds a note and a mood in one tap, so managers get a daily summary without chasing anyone.'], 17, False, DARK)

# 8 records
s = new('Records and My day', 'Transparent for the employee, ready for HR')
card(s, 0.6, 1.9, 5.9, 2.3, 'Records', ['Month view: days present, hours, overtime, late days', 'Leave balance: Casual, Sick, Earned', 'My requests with status', 'Missed-punch fixes'], TEAL, 17, 13)
card(s, 6.8, 1.9, 5.9, 2.3, 'My day', ['Every punch with time and place', 'Breaks between visits', 'Shift timing and today\'s notices'], BLUE, 17, 13)
tiles = [('Days present', TEAL, MINT), ('Hours worked', BLUE, RGBColor(0xE9, 0xF1, 0xFC)), ('Overtime', PURPLE, RGBColor(0xF0, 0xEA, 0xFB)), ('Late days', AMBER, RGBColor(0xFD, 0xF0, 0xE3))]
for i, (t, c, f) in enumerate(tiles):
    box(s, 0.6 + i * 3.1, 4.8, 2.9, 1.5, f)
    text(s, 0.6 + i * 3.1, 5.1, 2.9, 0.9, t, 20, True, c, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)

# 9 web
s = new('The web side for HR and managers', 'Four screens in the HR module')
card(s, 0.6, 1.9, 5.9, 2.2, 'Attendance requests', ['Approve or reject WFH, leave, late, corrections', 'Hear the voice note, add a decision note', 'Managers see their team; admin sees all'], TEAL, 17, 13)
card(s, 6.8, 1.9, 5.9, 2.2, 'Announcements', ['Title, body, tag, dates', 'Require acknowledgement and see who read it', 'Hide or show any time'], BLUE, 17, 13)
card(s, 0.6, 4.4, 5.9, 2.2, 'App Home messages', ['Control the daily heading and the thought of the day', 'Add many; one is chosen per day', 'No app update needed'], PURPLE, 17, 13)
card(s, 6.8, 4.4, 5.9, 2.2, 'Visit plans', ['Pick date, employee, customer (from your customer list), time and task', 'See which stops are done', 'Shows up in the app as Today\'s plan'], AMBER, 17, 13)

# 10 location privacy
s = new('Location and privacy, handled openly', 'Only what is needed, only while it is needed')
card(s, 0.6, 1.9, 5.9, 2.3, 'Geofence', ['Each site has a radius', 'Leaving and re-entering is recorded with the time', 'The app and Day complete show total time outside'], TEAL, 17, 13)
card(s, 6.8, 1.9, 5.9, 2.3, 'Tracking window', ['Location is tracked only until check-out', 'A clear note on Home explains it', 'Tap it to see "What we collect"'], BLUE, 17, 13)
card(s, 0.6, 4.5, 5.9, 2.0, 'Face data', 'Only a numeric face signature is used for matching, never a video. Fallback exists if a face is not recognised.', PURPLE, 17, 13)
card(s, 6.8, 4.5, 5.9, 2.0, 'Security', 'Signed login token, admin-only actions on the web, protected forms, inputs sanitised.', AMBER, 17, 13)

# 11 reminders & languages
s = new('Gentle reminders, in your language', 'Helpful, capped, never spammy')
card(s, 0.6, 1.9, 3.9, 2.8, 'Reminders', ['One check-in reminder per working day', 'Approval / rejection alerts', 'Managers: requests waiting', 'Urgent announcements', '"Looks like you left" check-out nudge'], TEAL, 17, 13)
card(s, 4.7, 1.9, 3.9, 2.8, 'Languages', ['English', 'मराठी', 'हिंदी', 'ગુજરાતી', 'Switch any time in Profile'], BLUE, 17, 13)
card(s, 8.8, 1.9, 3.9, 2.8, 'Consistent design', ['Same top bar and bottom bar on every screen', 'Clean cards, large touch targets', 'Fits a small phone without clutter'], PURPLE, 17, 13)
box(s, 0.6, 5.0, 12.1, 1.5, MINT)
text(s, 0.9, 5.2, 11.5, 1.2, 'Notifications are local and arrive within about 30 minutes. Real push (FCM) is on the roadmap.', 17, False, DARK)

# 12 how to use
s = new('Quick guide: how to use it', 'Three roles, three short paths')
card(s, 0.6, 1.9, 3.9, 4.6, 'Employee (daily)', ['1. Open the app', '2. Tap Check In, look at the camera', '3. Use the tiles if you will be late / away', '4. Add notes / photos during visits', '5. End visit or check out, add a day note', '6. Check Records for hours and leave'], TEAL, 17, 13)
card(s, 4.7, 1.9, 3.9, 4.6, 'Manager', ['1. Open Attendance requests on the web', '2. Review pending items (with voice notes)', '3. Approve or reject with a note', '4. Check team attendance and corrections', '5. Acknowledge face-less check-ins'], BLUE, 17, 13)
card(s, 8.8, 1.9, 3.9, 4.6, 'Admin / HR', ['1. Add announcements', '2. Set the daily heading and thoughts', '3. Plan visits for the field team', '4. Maintain sites, shifts and leave balances', '5. Review reports'], PURPLE, 17, 13)

# 13 flow
s = new('How the pieces talk to each other')
cols = [('Android app', ['Home, Capture, Records,', 'My day, Profile', 'Local alarms and tracking'], TEAL),
        ('Attendance API', ['Token login', 'Check-in / out, face match', 'Requests, visits, notes,', 'announcements, records'], BLUE),
        ('Database', ['Attendance, faces', 'Requests, balances, shifts', 'Announcements, messages', 'Visit plans and notes'], PURPLE),
        ('HR web', ['Approve / reject', 'Announcements, messages', 'Visit plans'], AMBER)]
for i, (h, b, c) in enumerate(cols):
    x = 0.6 + i * 3.15
    box(s, x, 2.0, 2.85, 3.3, WHITE, c)
    box(s, x, 2.0, 2.85, 0.6, c)
    text(s, x, 2.05, 2.85, 0.5, h, 17, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    text(s, x + 0.15, 2.8, 2.6, 2.4, b, 13, False, DARK)
    if i < 3:
        a = s.shapes.add_shape(MSO_SHAPE.LEFT_RIGHT_ARROW, Inches(x + 2.86), Inches(3.4), Inches(0.28), Inches(0.4))
        a.fill.solid(); a.fill.fore_color.rgb = TEAL; a.line.fill.background()
box(s, 0.6, 5.7, 12.1, 0.9, MINT)
text(s, 0.9, 5.85, 11.5, 0.7, 'Example: Employee requests leave  ->  API saves it  ->  Manager approves on web  ->  Balance updates  ->  Employee is notified', 15, True, TEAL)

# 14 roadmap
s = new('What comes next', 'Planned improvements')
card(s, 0.6, 1.9, 3.9, 3.6, 'Soon', ['Real push notifications', 'Widget and lock-screen check-in', 'Name an unknown site in the app', 'Announcements per branch / team'], TEAL, 17, 13)
card(s, 4.7, 1.9, 3.9, 3.6, 'Next', ['Weather and festival cards', 'Badges and streak rewards', 'Dark mode', 'Travel allowance and kudos'], BLUE, 17, 13)
card(s, 8.8, 1.9, 3.9, 3.6, 'Later', ['Planned visit routes and km', 'Customer insights on the visit card', 'Wearable support', 'Advanced analytics'], PURPLE, 17, 13)

# 15 close
n += 1
s = prs.slides.add_slide(blank)
bg(s, TEAL)
text(s, 0.9, 2.5, 11.5, 1.2, 'Thank you', 54, True, WHITE)
text(s, 0.9, 3.8, 11.5, 1.0, 'Questions and live demo', 24, False, RGBColor(0xCF, 0xE8, 0xE0))

out = r'C:\wamp64\www\captured_attendance\docs\ElintOmAT_Product_Overview.pptx'
prs.save(out)
print('saved', out, 'slides', len(prs.slides))
