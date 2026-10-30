from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

TEAL = RGBColor(0x0E, 0x6B, 0x5C)
doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
for m in ('left_margin', 'right_margin'):
    setattr(sec, m, Cm(1.6))
sec.top_margin = sec.bottom_margin = Cm(1.5)

st = doc.styles['Normal']
st.font.name = 'Calibri'
st.font.size = Pt(10)
for name, size in (('Heading 1', 18), ('Heading 2', 13), ('Heading 3', 11)):
    h = doc.styles[name]
    h.font.name = 'Calibri'
    h.font.size = Pt(size)
    h.font.color.rgb = TEAL
    h.font.bold = True


def shade(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear')
    sh.set(qn('w:color'), 'auto')
    sh.set(qn('w:fill'), hex_)
    tcPr.append(sh)


def para(text, bold=False, size=None, color=None, italic=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    return p


def bullets(items):
    for i in items:
        doc.add_paragraph(i, style='List Bullet')


def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.autofit = False
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ''
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(255, 255, 255)
        shade(c, '0E6B5C')
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ''
            run = cells[i].paragraphs[0].add_run(str(v))
            run.font.size = Pt(8.5)
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
    # repeat header row
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement('w:tblHeader')
    th.set(qn('w:val'), 'true')
    trPr.append(th)
    doc.add_paragraph()
    return t


# ------------------------------------------------------------------ cover
p = para('ElintOmAT', bold=True, size=30, color=TEAL)
para('Attendance & Work Companion – QA Test Plan and Test Cases', bold=True, size=16)
para('Android app  +  Web (HR / Manager) + API', size=12)
para('Version 1.0   |   Date: 30 Oct 2026   |   Audience: QA team', size=10, italic=True)
doc.add_paragraph()

doc.add_heading('1. Purpose and scope', 1)
para('This document lists everything the QA team must verify for the ElintOmAT release: the Android employee app, the web screens '
     'used by HR / managers, and the API between them. Each test case has an ID, steps, expected result and priority. '
     'Run the cases in the order of the sections; the end-to-end scenarios in section 8 depend on the earlier sections.')
doc.add_heading('In scope', 2)
bullets([
    'Android app (package com.aistudio.captureattendance.vyzp): login, face enrolment, Home, check-in / check-out, quick updates, field-visit tools, Records, My Day, Profile, languages, reminders, geofence tracking.',
    'Web (HR module): Attendance requests (approve / reject), Announcements, App Home messages, Visit plans.',
    'API: attendance_api endpoints used by the app (authentication, validation, error handling).',
    'Cross-cutting: offline behaviour, security, localisation (EN / MR / HI / GU), performance, device compatibility.',
])
doc.add_heading('Out of scope for this release (known gaps – do not log as defects)', 2)
bullets([
    'Push notifications through FCM (the app uses local alarms + 30-minute polling instead).',
    'Home-screen widget, lock-screen tile, smartwatch.',
    'Naming UI for an unknown site (an unknown site is saved automatically but cannot be renamed in the app).',
    'Announcement targeting by branch / team / role (all announcements go to everyone).',
    'Dark mode, badges, travel allowance, kudos, AI avatar, weather, km travelled (roadmap Phase 2 / 3).',
])

doc.add_heading('2. Test environment', 1)
table(['Item', 'Details'], [
    ['Android devices', 'Minimum 3 real devices: Android 9–10 (low-end), Android 12–13 (mid), Android 14+ (latest). Screen sizes 5.5", 6.1", 6.7". At least one with a notch / punch-hole camera.'],
    ['Network profiles', 'Wi-Fi, 4G, 2G-like throttled, flight mode (offline), airplane mode toggled during an action.'],
    ['Backend', 'ElintOm_PHP_8.5 (CodeIgniter, PHP 8.5, MySQL). Tenant DB with prefix sma_. API base: <server>/attendance_api/<endpoint>.'],
    ['Web browsers', 'Chrome (latest), Edge (latest), Firefox (latest). Desktop 1366x768 and 1920x1080.'],
    ['Accounts', 'A) Admin (full access)  B) Employee with a reporting manager  C) Employee without a reporting manager  D) Manager (has direct reports)  E) Employee with no face registered.'],
    ['Master data', 'At least 2 saved sites with latitude / longitude / radius; 5+ customers (sma_companies group "customer"); shift 09:30–18:30 (default); leave balance Casual 12, Sick 12, Earned 15.'],
    ['Tools', 'Android Studio logcat, Postman / curl, browser dev-tools, a GPS-spoofing app (for geofence tests on a test device only).'],
], [5, 20])

doc.add_heading('3. Entry / exit criteria and priorities', 1)
bullets([
    'Entry: build installed, backend reachable, test accounts and master data created, tester has read sections 1–2.',
    'Exit: 100% of P1 cases pass, at least 95% of P2 pass, no open Critical or High defect, all failed P3 cases have an agreed workaround.',
    'Priority: P1 = release blocker (core attendance, security, data loss, crash). P2 = important feature. P3 = cosmetic / minor.',
    'Severity: Critical (crash, data loss, security), High (feature unusable), Medium (workaround exists), Low (cosmetic).',
])

# ------------------------------------------------------------------ test case data
H = ['ID', 'Title', 'Steps', 'Expected result', 'Pri']
W = [2.0, 5.0, 9.0, 8.5, 1.0]

sections = []

sections.append(('4.1 Android – Login, setup and permissions', [
    ['AND-LOG-01', 'Login with valid credentials', '1) Install app and open. 2) Enter valid user and password. 3) Tap Login.', 'Login succeeds, Home opens, name shown in top bar, token saved (user stays logged in after app restart).', 'P1'],
    ['AND-LOG-02', 'Login with wrong password', 'Enter valid user and wrong password.', 'Clear error message; no crash; fields stay filled; can retry.', 'P1'],
    ['AND-LOG-03', 'Login without network', 'Switch on flight mode, try to log in.', 'Friendly "no connection" message; no crash or endless spinner.', 'P2'],
    ['AND-LOG-04', 'Session expiry / 401', 'Invalidate the token on the server (or wait for expiry) and open Home.', 'App logs the user out and shows Login; no crash.', 'P1'],
    ['AND-LOG-05', 'Logout', 'Profile > Log out > confirm.', 'Returns to Login; reminders and area tracking stop; back button does not return to Home.', 'P1'],
    ['AND-LOG-06', 'Permissions card', 'Profile > Permissions. Deny camera, then location, then notifications.', 'Card shows count (e.g. 4/6) and a working Allow / Open settings action for each missing permission.', 'P2'],
    ['AND-LOG-07', 'Notification permission prompt (Android 13+)', 'Fresh install, login, open Home.', 'Permission asked once, not on every launch.', 'P2'],
    ['AND-LOG-08', 'Face enrolment', 'Profile > Register / Update Face, follow on-screen guide, submit.', 'Face saved; success message; later check-in recognises the user.', 'P1'],
    ['AND-LOG-09', 'Face enrolment – no camera permission', 'Deny camera, open enrolment.', 'Prompt to grant permission; no crash.', 'P2'],
]))

sections.append(('4.2 Android – App shell, top bar and navigation', [
    ['AND-NAV-01', 'Common top bar', 'Open Home, My day, Records, Profile in turn.', 'Same full-width white top bar (avatar, greeting, name, date, bell, calendar) on all four; Capture screen has its own back header.', 'P2'],
    ['AND-NAV-02', 'Bottom navigation', 'Tap Home, My day, Records, Profile.', 'Correct screen opens; selected tab highlighted in teal; nav bar fixed at bottom.', 'P1'],
    ['AND-NAV-03', 'Avatar tap', 'Tap avatar in top bar.', 'Opens Profile.', 'P3'],
    ['AND-NAV-04', 'Calendar icon', 'Tap calendar icon.', 'Opens Records (new themed screen, not the old Attendance History).', 'P2'],
    ['AND-NAV-05', 'Bell icon', 'Tap bell icon.', 'Announcements list opens; red dot only when unread items exist.', 'P2'],
    ['AND-NAV-06', 'Back button', 'From Capture / Records / Profile press device back.', 'Returns to Home; from Home exits the app (no loop).', 'P2'],
    ['AND-NAV-07', 'Status bar / insets', 'Open every tab on a notch device and on a gesture-nav device.', 'Content is not hidden behind status bar or nav bar.', 'P2'],
]))

sections.append(('4.3 Android – Home (not checked in)', [
    ['AND-HOME-01', 'Layout fits without empty space', 'Open Home on 5.5", 6.1", 6.7" devices.', 'Content fills the screen with compact, even spacing; no large empty area at the bottom; scrolls only if content is taller than the screen.', 'P2'],
    ['AND-HOME-02', 'Daily line is dynamic', 'On web add / hide lines in App Home messages (type: Daily line). Pull Home again.', 'Home shows the line chosen for today from active lines; hidden lines never appear; if none active the app default is shown.', 'P2'],
    ['AND-HOME-03', "Today's thought is dynamic", 'On web add / hide Thoughts. Reopen Home.', "Card shows today's active thought; heart toggles favourite and survives restart; no share button.", 'P2'],
    ['AND-HOME-04', 'Check-in card', 'Observe card when not checked in.', 'Shows "Not Checked In", shift time, large Check In button, location pill, recognised-site badge when matched, View on Map link.', 'P1'],
    ['AND-HOME-05', 'Location shows live position', 'Move to a different place (or spoof GPS), reopen Home.', 'Pill shows "Fetching location…" then the current landmark / coordinates; never an old location from earlier.', 'P1'],
    ['AND-HOME-06', 'View on Map', 'Tap View on Map.', 'Maps app opens with a pin at the current live position.', 'P2'],
    ['AND-HOME-07', 'Location pill tap', 'Tap the location pill.', 'Fresh GPS fix is requested and a detail dialog opens.', 'P3'],
    ['AND-HOME-08', 'Stats row', 'Check stats after several days of data.', '"x of 5 days this week", on-time streak (10-min grace), month % match Records and server data.', 'P2'],
    ['AND-HOME-09', 'Stats tap', 'Tap a stats card.', 'Opens Records.', 'P3'],
    ['AND-HOME-10', 'Announcement strip – read more', 'Create a short and a long announcement.', 'Short text: no "Read more". Long text that overflows one line: "Read more" shown; tap expands, tap text collapses.', 'P2'],
    ['AND-HOME-11', 'Announcement pager', 'Create 3 announcements. Swipe left/right and use arrows.', 'Pager cycles with n/total counter; needs-acknowledge items are shown first.', 'P2'],
    ['AND-HOME-12', 'Acknowledge announcement', 'Tap Acknowledge on a needs-ack item.', 'Item marked acknowledged; button disappears; web read-receipt count increases by 1.', 'P2'],
    ['AND-HOME-13', 'Tracking note', 'Tap the "Location is tracked only until you check out" note.', '"What we collect" dialog opens with correct wording.', 'P3'],
    ['AND-HOME-14', 'Pull / reopen refresh', 'Change data on the server, switch tabs and return.', 'Home refreshes summary data (state, announcements, thought).', 'P2'],
]))

sections.append(('4.4 Android – Quick updates (Running late / WFH / Can\'t come / Take leave)', [
    ['AND-QU-01', 'Open sheet', 'Tap each of the four tiles.', 'Update sheet opens with the correct chip selected; X closes it.', 'P1'],
    ['AND-QU-02', 'Running late', 'Choose 15 / 30 / 45 min / 1 h, optional note, Send.', 'Request status "sent"; line on Home reads e.g. "Running late – expected 10:00 AM"; manager sees request on web.', 'P1'],
    ['AND-QU-03', 'Work from home', 'Choose Today / Tomorrow, duration (Full / First half / Second half), Send.', 'Status line on Home; info box says manager approval needed; web shows pending WFH.', 'P1'],
    ['AND-QU-04', "Can't come", 'Pick a reason (Unwell, Family, Travel, Other) and Send.', 'Request sent; Home line shows the reason.', 'P1'],
    ['AND-QU-05', 'Take leave – date range', 'Pick Casual / Sick / Earned, open date range picker, select dates, Send.', 'Range shown on button; request created with correct from/to dates and working-day count.', 'P1'],
    ['AND-QU-06', 'Take leave – balance check', 'Request more days than the remaining balance.', 'Clear error or warning; no request created beyond balance (per server rule).', 'P2'],
    ['AND-QU-07', 'Undo', 'Tap Undo on a sent / pending request.', 'Request removed from Home; web status becomes undone and disappears from the manager inbox.', 'P1'],
    ['AND-QU-08', 'Approved state', 'Approve the request on web, reopen Home.', 'Green line "Approved"; X dismisses it and it does not return.', 'P1'],
    ['AND-QU-09', 'Rejected state', 'Reject the request with a note on web.', 'Red line; View shows the manager note.', 'P1'],
    ['AND-QU-10', 'Voice note on request', 'Tap mic, allow permission, record, stop, Send.', 'Request is created with the audio; web shows an audio player that plays the note; max length enforced.', 'P2'],
    ['AND-QU-11', 'Mic permission denied', 'Deny microphone permission.', 'Message "microphone needed"; sheet still usable without voice.', 'P2'],
    ['AND-QU-12', 'No reporting manager', 'Login as account C and send a request.', 'Request saved; message explains no manager is assigned (per server wording); no crash.', 'P2'],
    ['AND-QU-13', 'Late / can\'t come auto-clear', 'Send "running late", then check in.', 'The late notice clears from Home after check-in.', 'P2'],
    ['AND-QU-14', 'Sending while offline', 'Flight mode, tap Send.', 'Error shown; sheet stays open with data kept; no duplicate request when retried.', 'P2'],
]))

sections.append(('4.5 Android – Check-in / check-out (Capture)', [
    ['AND-CAP-01', 'Open Capture', 'Tap Check In on Home.', 'Capture screen with back arrow, location card, IN/OUT toggle, last activity, notes, camera frame with face oval, "Capture Attendance" button.', 'P1'],
    ['AND-CAP-02', 'Successful check-in', 'Face the camera, tap Capture.', 'Success toast + vibration; app returns to Home automatically; Home shows Checked in header and Visit card; server has a new row in attendance.', 'P1'],
    ['AND-CAP-03', 'Successful check-out', 'From Home tap End visit, scan face.', 'Success; Home shows Day complete; server row has check_out time.', 'P1'],
    ['AND-CAP-04', 'Face not recognised', 'Use a different person / cover the face.', '"Face not recognised" error; user stays on Capture and can retry; no attendance saved.', 'P1'],
    ['AND-CAP-05', 'Non-biometric offer after 3 failures', 'Fail the face match 3 times in a row.', 'Dialog offers "Check in without face scan"; confirming records the punch and creates a "manual" request for the manager.', 'P1'],
    ['AND-CAP-06', 'Non-biometric – decline', 'Dismiss the dialog.', 'No attendance saved; counter resets; user can continue scanning.', 'P2'],
    ['AND-CAP-07', 'Non-biometric – web visibility', 'After AND-CAP-05 open Attendance requests on web.', 'A request "Check-in without face scan" shows the user, time and place; system_log of the punch contains "non_biometric".', 'P1'],
    ['AND-CAP-08', 'No face detected', 'Point camera at a wall.', 'Message "No face detected" and guide stays; Capture does not submit.', 'P2'],
    ['AND-CAP-09', 'GPS not ready', 'Disable GPS or deny location, tap Capture.', 'Clear message / prompt to enable location; no crash; no punch saved.', 'P1'],
    ['AND-CAP-10', 'Low GPS accuracy', 'Test indoors with poor GPS.', 'Accuracy label "coarse" shown; behaviour follows server rule; tapping the location card refreshes.', 'P2'],
    ['AND-CAP-11', 'Site matching', 'Check in inside a saved site radius, then outside any site.', 'Inside: site name and recognised badge. Outside: allowed; location saved as a new site (flagged "new location found").', 'P1'],
    ['AND-CAP-12', 'Approved WFH day', 'Approve a WFH for today, then check in.', 'Punch location is "WFH".', 'P1'],
    ['AND-CAP-13', 'Notes with punch', 'Type a note on Capture and submit.', 'Note stored with the punch and cleared afterwards.', 'P3'],
    ['AND-CAP-14', 'Double tap protection', 'Tap Capture rapidly twice.', 'Only one punch is created.', 'P1'],
    ['AND-CAP-15', 'Duplicate check-in guard', 'Try to check in again while already checked in (use IN toggle).', 'Server rejects or switches to OUT according to rule; no duplicate open row.', 'P1'],
    ['AND-CAP-16', 'Check-out without check-in', 'Force OUT with no open check-in.', 'Error "No active check-in found"; no row created.', 'P2'],
    ['AND-CAP-17', 'Auto Submit', 'Profile > Auto Submit ON, open Capture.', 'Capture submits automatically when a good face frame is found; OFF requires tapping the button.', 'P2'],
    ['AND-CAP-18', 'Camera permission denied', 'Deny camera and open Capture.', 'Prompt with Allow action; no crash.', 'P1'],
    ['AND-CAP-19', 'Back from Capture', 'Press the back arrow on Capture.', 'Returns to Home; camera released.', 'P2'],
]))

sections.append(('4.6 Android – Offline punches and sync', [
    ['AND-OFF-01', 'Check-in offline', 'Flight mode, complete a check-in.', 'Toast "saved offline"; Home shows a "Pending sync" note; Home still shows Checked in state.', 'P1'],
    ['AND-OFF-02', 'Sync after reconnect', 'Turn the network back on.', 'Queued punch uploads automatically; "Synced" toast; Pending note disappears; server row has the original punch time.', 'P1'],
    ['AND-OFF-03', 'Multiple offline punches', 'Make check-in and check-out offline, reconnect.', 'Both sync in order; one rejected punch does not block the other.', 'P2'],
    ['AND-OFF-04', 'App killed while pending', 'Create an offline punch, force-stop the app, reopen online.', 'Queue survives and is uploaded.', 'P1'],
    ['AND-OFF-05', 'Offline Home', 'Open Home in flight mode.', 'Last known data shown or graceful empty state; no crash.', 'P2'],
]))

sections.append(('4.7 Android – Checked-in Home, timer and visit tools', [
    ['AND-VIS-01', 'Checked-in header', 'After check-in open Home.', 'Top bar shows green dot "Checked in · Xh Ym" and "Keep going, <first name>!".', 'P2'],
    ['AND-VIS-02', 'Visit timer starts at check-in', 'Check in and watch the timer.', 'Timer starts from the check-in moment and counts every second.', 'P1'],
    ['AND-VIS-03', 'Timer does not reset', 'Switch tabs, lock / unlock phone, force-stop and reopen the app.', 'Timer continues from the original check-in time and never returns to 00:00:00.', 'P1'],
    ['AND-VIS-04', 'Timer cleared on check-out', 'Check out and check in again.', 'New timer starts from the new check-in.', 'P1'],
    ['AND-VIS-05', 'Visit card site name', 'Check in at a saved site.', 'Card shows the planned stop name (if a plan exists) or the punch location; never "Acquiring GPS location…" while checked in.', 'P2'],
    ['AND-VIS-06', 'Visit card stats', 'Create a plan for a customer with open orders and an earlier done visit.', 'Chips show open orders, last visit date, task due count that match the data.', 'P2'],
    ['AND-VIS-07', 'Add note', 'Tap Add note, type text, Save.', 'No crash; "Saved" toast; note appears in Details and in server table visit_notes.', 'P1'],
    ['AND-VIS-08', 'Add note – empty', 'Open Add note and leave it blank.', 'Save is disabled.', 'P3'],
    ['AND-VIS-09', 'Photo', 'Tap Photo, allow camera, take a picture.', 'Photo uploads; "Saved" toast; server file saved under attendance_visit and listed in Details.', 'P2'],
    ['AND-VIS-10', 'Voice note', 'Tap Voice note, allow mic, record, stop, Save.', 'Audio uploads (max 60 s); "Saved" toast; web / Details list it.', 'P2'],
    ['AND-VIS-11', 'Details dialog', 'Tap Details.', 'Customer name, address, phone, task, and today\'s notes (with time) are listed.', 'P2'],
    ['AND-VIS-12', "Today's plan list", 'Create 6 plan items for the user on web.', 'Card lists items in time order with Done ticks, one "Now", others empty circles; header "x of 6 done"; "See all 6 visits" expands to all and "Show less" collapses.', 'P2'],
    ['AND-VIS-13', 'No plan', 'User with no plan today.', "Today's plan card is not shown; visit card shows the site name.", 'P2'],
    ['AND-VIS-14', 'Plan advances on End visit', 'Tap End visit and check out.', 'Current stop becomes Done on the server; next stop becomes Now when checking in again.', 'P1'],
    ['AND-VIS-15', 'Plan belongs to user', 'Assign a plan to user A, login as user B.', 'User B does not see user A\'s plan.', 'P1'],
    ['AND-VIS-16', 'Outside-area line', 'Walk outside the site radius and back (see AND-GEO tests), reopen Home.', 'Visit card shows "Outside the area: Xh Ym (n x)" in yellow; hidden when never outside.', 'P2'],
]))

sections.append(('4.8 Android – Day complete', [
    ['AND-DAY-01', 'Day complete screen', 'Check out and open Home.', 'Shows check-out time and place (cleaned text, not raw "Check In : … | Check Out : …"), worked time, number of visits.', 'P1'],
    ['AND-DAY-02', 'Save note and mood', 'Type a note, choose a mood, Save.', 'No crash; "Saved" toast; values persist after reopening the app; server row in attendance_day_notes.', 'P1'],
    ['AND-DAY-03', 'Outside-area summary', 'After a day with an exit, view Day complete.', 'Shows the outside-area total and number of exits.', 'P2'],
    ['AND-DAY-04', 'Save without network', 'Go offline and Save.', 'Error toast; text stays in the box.', 'P2'],
    ['AND-DAY-05', 'Second check-in same day', 'Check in again after checking out.', 'Home returns to the checked-in view; visits count increases.', 'P2'],
]))

sections.append(('4.9 Android – Records and My Day', [
    ['AND-REC-01', 'Records layout', 'Open Records.', 'Month bar with arrows, four tinted tiles (Days present, Hours worked, Overtime, Late days), Leave balance, My requests, attendance list.', 'P2'],
    ['AND-REC-02', 'Month navigation', 'Tap previous / next month.', 'Data and totals reload for the chosen month; no crash at year boundaries.', 'P2'],
    ['AND-REC-03', 'Totals accuracy', 'Compare tiles with the raw attendance rows of the month.', 'Days present, worked minutes, overtime beyond shift, late days (after shift start + grace) are correct.', 'P1'],
    ['AND-REC-04', 'Leave balance', 'Approve a casual leave of 2 days on web.', 'Casual remaining decreases by 2 (shown as remaining / allotted).', 'P1'],
    ['AND-REC-05', 'My requests list', 'Create late, WFH, leave requests.', 'Each listed with kind, dates and status (sent, pending, approved, rejected); undone ones hidden.', 'P2'],
    ['AND-REC-06', 'Empty states', 'New user with no data.', '"No requests yet." with icon and "No attendance in this month." card.', 'P3'],
    ['AND-REC-07', 'Missed punch card', 'Leave an open check-in from a previous day.', 'Card suggests the shift-end time; "Request correction" sends a correction (10-second flow).', 'P1'],
    ['AND-REC-08', 'Correction validations', 'Send duplicate correction / time earlier than check-in.', 'Server rejects with a clear message.', 'P2'],
    ['AND-REC-09', 'Correction approval', 'Approve the correction on web.', 'Check-out is applied to the attendance row; missed-punch card disappears.', 'P1'],
    ['AND-MYD-01', 'My Day timeline', 'Open My day after 2 punches.', 'Punches in order with times, break between check-out and next check-in, shift info.', 'P2'],
    ['AND-MYD-02', 'My Day with request', 'Create a request and open My day.', 'Request notice shown with status.', 'P3'],
]))

sections.append(('4.10 Android – Profile, language and settings', [
    ['AND-PRO-01', 'Profile layout', 'Open Profile.', 'Common top bar, cards for Register Face, Permissions, Language, Auto Submit, System & Diagnostics, red Log out; no duplicate name header; same theme as Home.', 'P2'],
    ['AND-PRO-02', 'Language switch', 'Choose English, Marathi, Hindi, Gujarati.', 'All screens change language immediately; no clipped or overlapping text; fonts render Devanagari / Gujarati correctly.', 'P1'],
    ['AND-PRO-03', 'Language persistence', 'Change language, restart.', 'Language persists.', 'P2'],
    ['AND-PRO-04', 'Auto Submit toggle', 'Switch on and off.', 'State persists and subtitle text changes.', 'P3'],
    ['AND-PRO-05', 'Diagnostics logs', 'Open System & Diagnostics > View logs.', 'Logs dialog opens; no sensitive token shown in plain text.', 'P3'],
    ['AND-PRO-06', 'Admin-only item', 'Login as admin vs employee.', '"Configure server URL" only visible to admin.', 'P2'],
]))

sections.append(('4.11 Android – Reminders, notifications and background tracking', [
    ['AND-REM-01', 'Daily check-in reminder', 'Not checked in, shift 09:30 (grace 10 min). Wait until 09:40 on a weekday.', 'One local notification; not repeated the same day; none on Sat/Sun.', 'P2'],
    ['AND-REM-02', 'No reminder when already told manager', 'Send "Running late" / WFH before 09:40.', 'No check-in reminder is shown.', 'P2'],
    ['AND-REM-03', 'Approval notification', 'Approve / reject a request on web; wait up to 30 minutes.', 'Employee gets a notification once (approved / rejected, with note).', 'P2'],
    ['AND-REM-04', 'Manager inbox notification', 'Employee submits request for a manager account.', 'Manager gets "n requests waiting" notification once per new request.', 'P2'],
    ['AND-REM-05', 'Urgent announcement notification', 'Create a needs-acknowledge announcement.', 'Notification appears once; tapping opens the app.', 'P3'],
    ['AND-REM-06', 'Quiet hours', 'Check behaviour between 21:00 and 07:00.', 'Polling alarm does not fire notifications at night.', 'P3'],
    ['AND-REM-07', 'Reminders stop on logout', 'Logout and wait for an alarm.', 'No notifications.', 'P2'],
    ['AND-REM-08', 'Notification permission denied', 'Deny notifications.', 'No crash; no notifications.', 'P3'],
    ['AND-GEO-01', 'Tracking starts on check-in', 'Check in at a site with a radius.', 'Foreground notification / service starts; Home shows tracking note.', 'P1'],
    ['AND-GEO-02', 'Exit event', 'Move outside the site radius (spoof GPS).', 'Server table attendance_area_log gets an "outside" row with distance and radius; "left the site" nudge shown.', 'P1'],
    ['AND-GEO-03', 'Re-entry event', 'Return inside the radius.', 'An "inside" row is added; no duplicate rows while staying inside / outside.', 'P1'],
    ['AND-GEO-04', 'Outside total calculation', 'Stay outside for ~10 minutes then return.', 'Home / Day complete show about 10 min and 1 exit; an exit still open at check-out is counted to check-out time.', 'P1'],
    ['AND-GEO-05', 'Tracking stops on check-out', 'Check out.', 'Service stops; tracking notification removed.', 'P1'],
    ['AND-GEO-06', 'Battery optimisation', 'Enable battery saver and doze for a few hours.', 'Tracking still logs exits (or the app guides the user to the settings).', 'P2'],
]))

sections.append(('5. Web (HR / Manager) test cases', [
    ['WEB-ACC-01', 'Menu entries', 'Login as admin; open the HR menu.', 'Entries visible: Attendance requests, App Home messages, Visit plans.', 'P2'],
    ['WEB-ACC-02', 'Non-admin access', 'Login as a normal user and open the three URLs.', 'Admin-only actions (add, delete, toggle, decide) are blocked; page shows access denied or read-only.', 'P1'],
    ['WEB-ACC-03', 'CSRF / method check', 'Send POST without CSRF token or use GET on save URLs.', 'Request is rejected; no data change.', 'P1'],
    ['WEB-REQ-01', 'List requests', 'Open Attendance requests.', 'Table lists requests with employee, kind, dates, duration, reason, status and voice player if any; filters (pending, approved, rejected, sent) work.', 'P1'],
    ['WEB-REQ-02', 'Approve WFH', 'Approve a pending WFH request with a note.', 'Status approved, decided_by / at saved; employee Home shows approved line; punch that day is marked WFH.', 'P1'],
    ['WEB-REQ-03', 'Approve leave', 'Approve a 2-day leave.', 'Leave balance reduces by working days; shown in Records.', 'P1'],
    ['WEB-REQ-04', 'Reject with note', 'Reject a request with a note.', 'Employee sees rejected + note; balance unchanged.', 'P1'],
    ['WEB-REQ-05', 'Already decided', 'Try to decide the same request twice (two tabs).', 'Second attempt is refused or ignored; no double deduction.', 'P1'],
    ['WEB-REQ-06', 'Approve correction', 'Approve a missed-punch correction.', 'check_out applied to the attendance row.', 'P1'],
    ['WEB-REQ-07', 'Manual (non-biometric) entry', 'Open a "Check-in without face scan" item.', 'Shows user, time, place; can be acknowledged / decided.', 'P2'],
    ['WEB-REQ-08', 'Manager scope', 'Login as a manager (not admin).', 'Sees only requests routed to them; admin sees all.', 'P1'],
    ['WEB-REQ-09', 'Voice note playback', 'Play audio on a request with voice.', 'Audio plays in browser; file path is under assets/uploads/attendance_voice.', 'P2'],
    ['WEB-ANN-01', 'Create announcement', 'Fill title, body, tag, needs-ack, start / end, Save.', 'Appears in list and in the app pager within a refresh; validation requires a title.', 'P1'],
    ['WEB-ANN-02', 'Date window', 'Set start in the future / end in the past.', 'Announcement shows in the app only inside its window.', 'P2'],
    ['WEB-ANN-03', 'Hide / show', 'Toggle an announcement.', 'Hidden one disappears from the app; shown again on re-enable.', 'P2'],
    ['WEB-ANN-04', 'Read receipts', 'Acknowledge from 2 users.', 'Acknowledged count shows 2.', 'P2'],
    ['WEB-ANN-05', 'XSS check', 'Enter <script>alert(1)</script> in title and body.', 'Text is escaped on web and in the app; no script runs.', 'P1'],
    ['WEB-MSG-01', 'Add daily line / thought', 'App Home messages > choose type, enter text, Save.', 'Row appears; max 255 characters; empty is rejected with a message.', 'P1'],
    ['WEB-MSG-02', 'Hide / delete', 'Hide a row; delete another.', 'Hidden row greyed and not sent to the app; deleted row removed.', 'P2'],
    ['WEB-MSG-03', 'Rotation', 'Have 3 active thoughts; check on consecutive days.', 'A different active thought is picked by day of year; same all day.', 'P3'],
    ['WEB-MSG-04', 'Default seed', 'Fresh database.', 'Default lines and thoughts exist after first load.', 'P3'],
    ['WEB-VIS-01', 'Add visit plan (customer)', 'Pick date, employee, customer, time, task; Save.', 'Row appears for that date; place and address auto-filled from the customer.', 'P1'],
    ['WEB-VIS-02', 'Add visit plan (free text)', 'Leave customer empty, enter place name.', 'Saved with the typed place.', 'P2'],
    ['WEB-VIS-03', 'Validation', 'Submit without employee or place.', 'Error "Choose an employee, a date and a customer or place"; nothing saved.', 'P2'],
    ['WEB-VIS-04', 'Change date', 'Use the date picker at the top.', 'List reloads for that date.', 'P3'],
    ['WEB-VIS-05', 'Delete plan', 'Delete a plan row and confirm.', 'Row removed; app no longer shows it.', 'P2'],
    ['WEB-VIS-06', 'Status updates', 'Employee checks out for a stop.', 'Row status becomes done.', 'P1'],
    ['WEB-ATT-01', 'Attendance details area log', 'Open an attendance record with outside/inside events.', 'Area movement log lists events with time, distance and radius.', 'P2'],
]))

sections.append(('6. API test cases (Postman / curl)', [
    ['API-01', 'Auth required', 'Call home_summary, my_day, records, quick_update, visit_context, visit_note, visit_details, manager_inbox without a token.', 'HTTP 401 with {"status":"error","code":"UNAUTHORIZED"}.', 'P1'],
    ['API-02', 'Tampered / expired token', 'Alter one character of the token; use an expired token.', 'HTTP 401.', 'P1'],
    ['API-03', 'home_summary', 'GET with a valid token.', 'JSON: user, shift, today (state, first_in, worked_minutes, outside_minutes, outside_count), request, announcements, stats, leave_balance, thought, hero, missed_punch.', 'P1'],
    ['API-04', 'store check-in / out', 'POST attendance_api/store with image descriptor and GPS.', 'Success with action, next_action, last_activity; row in attendance; invalid GPS returns an error.', 'P1'],
    ['API-05', 'store non_biometric', 'POST non_biometric=1 with and without a token.', 'With token: punch saved + manual request. Without token: UNAUTHORIZED.', 'P1'],
    ['API-06', 'quick_update validation', 'Send invalid kind, missing dates for leave, invalid duration.', 'HTTP 200 with status error and a message; nothing saved.', 'P1'],
    ['API-07', 'quick_update voice', 'Send a non-m4a / oversized base64.', 'Rejected with "Voice note is too long or not valid".', 'P2'],
    ['API-08', 'quick_update_undo', 'Undo own request; try another user\'s request id.', 'Own: undone. Other user: refused.', 'P1'],
    ['API-09', 'announcement_ack', 'Acknowledge twice.', 'Idempotent; one read row.', 'P3'],
    ['API-10', 'day_wrapup', 'POST note > 1000 chars, mood 0 / 6.', 'Trimmed / validated; mood outside 1–5 ignored.', 'P3'],
    ['API-11', 'records', 'GET records?month=2026-10 and an invalid month.', 'Valid month returns days and totals; invalid month falls back or returns an error, never a 500.', 'P2'],
    ['API-12', 'fix_missed_punch', 'Duplicate request; time before check-in; other user\'s attendance id.', 'Each rejected with a clear message.', 'P1'],
    ['API-13', 'manager_inbox', 'Call as manager and as employee.', 'Manager: pending_count and latest_id; employee: 0.', 'P2'],
    ['API-14', 'visit_context', 'GET with and without plan.', 'items[], done, total, current (open_orders, last_visit, tasks_due) or empty / null.', 'P1'],
    ['API-15', 'visit_note', 'POST kind=note/photo/voice; invalid kind; photo not JPEG; empty note.', 'Valid kinds saved; invalid rejected with message.', 'P1'],
    ['API-16', 'visit_details', 'GET after saving notes.', 'current, customer card and today\'s notes (kind, text, file url, time).', 'P2'],
    ['API-17', 'area_log', 'POST inside / outside coordinates for a site with radius.', 'Event logged only on change; response has event, distance_m, radius_m.', 'P1'],
    ['API-18', 'SQL injection / XSS', 'Send quotes, SQL fragments and HTML in text fields.', 'Parameters are escaped; no error or injected data.', 'P1'],
    ['API-19', 'Rate / load', 'Run 50 parallel home_summary calls.', 'Responses < 1.5 s average; no errors.', 'P3'],
]))

sections.append(('7. Non-functional test cases', [
    ['NFR-01', 'Cold start time', 'Start the app from cold on a mid device.', 'Home visible in under 3 seconds on Wi-Fi.', 'P2'],
    ['NFR-02', 'Home refresh time', 'Open Home on 4G.', 'Summary loaded in under 2 seconds.', 'P2'],
    ['NFR-03', 'Face scan time', 'Measure from camera open to punch result.', 'Under 5 seconds on a good network.', 'P2'],
    ['NFR-04', 'Battery usage', 'Stay checked in for 8 hours with tracking.', 'Battery drain within agreed limit (target < 8%/8h).', 'P2'],
    ['NFR-05', 'Rotation / config change', 'Rotate device on each screen.', 'State is kept; no crash.', 'P3'],
    ['NFR-06', 'Font scale / accessibility', 'Set system font to largest.', 'Text remains readable; buttons not clipped.', 'P3'],
    ['NFR-07', 'Memory', 'Switch tabs 50 times; use camera repeatedly.', 'No OutOfMemory or crash.', 'P2'],
    ['NFR-08', 'Data privacy', 'Inspect network traffic and logs.', 'Token and face data are not shown in logs; HTTPS used in production.', 'P1'],
    ['NFR-09', 'Face data', 'Check what is sent for a punch.', 'Only the 128-d descriptor (no raw video); stored in sma_user_faces.', 'P1'],
    ['NFR-10', 'Time zone', 'Use a device in a different time zone.', 'Punch times follow server time; Home / Records show consistent values.', 'P2'],
    ['NFR-11', 'Crash-free run', 'Run the full regression (sections 4–5) with logcat open.', 'No FATAL EXCEPTION. (Previously found: text-field NullPointer with ElintTypography – verify it is gone.)', 'P1'],
    ['NFR-12', 'Language completeness', 'Walk all screens in MR, HI and GU.', 'No English fallback strings or missing placeholders like {name}.', 'P2'],
    ['NFR-13', 'Web responsive', 'Open web screens at 1366 and 1920 width and in a narrow window.', 'Tables scroll, forms stay usable.', 'P3'],
    ['NFR-14', 'Browser matrix', 'Run web cases on Chrome, Edge, Firefox.', 'Same behaviour on all.', 'P3'],
]))

sections.append(('8. End-to-end scenarios', [
    ['E2E-01', 'WFH request lifecycle', '1) Employee sends WFH for today. 2) Manager sees it on web and approves. 3) Employee reopens Home. 4) Employee checks in.', 'Home shows Approved; punch is WFH; Records shows the day; manager inbox clears.', 'P1'],
    ['E2E-02', 'Leave lifecycle', '1) Employee requests 2 days casual leave. 2) Manager rejects first with note. 3) Employee re-requests. 4) Manager approves.', 'Rejection note visible; approval deducts 2 days; Records balance 10 / 12; notification received.', 'P1'],
    ['E2E-03', 'Missed check-out', '1) Employee checks in and forgets to check out. 2) Next day open Home. 3) Send correction. 4) Manager approves.', 'Missed-punch card shown, correction appears on web, approval writes the check-out, card disappears.', 'P1'],
    ['E2E-04', 'Field day', '1) Admin creates 3 plan items for an employee. 2) Employee checks in. 3) Adds note, photo, voice note. 4) Ends visit, checks in again for the next stop. 5) Checks out for the day.', 'Plan advances stop by stop; notes visible in Details; Day complete shows worked time and visits; web plan statuses done.', 'P1'],
    ['E2E-05', 'Face failure fallback', '1) Fail face match 3 times. 2) Accept non-biometric check-in. 3) Manager checks web.', 'Punch saved with non_biometric flag; manual request visible to manager.', 'P1'],
    ['E2E-06', 'Offline day', '1) Go offline, check in. 2) Add visit notes. 3) Reconnect.', 'Punch syncs with original time; Pending note disappears; (note: visit notes require connectivity – confirm expected behaviour).', 'P2'],
    ['E2E-07', 'Geofence day', '1) Check in at site. 2) Leave radius for 15 min and return. 3) Check out.', 'Area log has outside + inside rows; Day complete shows ~15 min outside, 1 exit.', 'P1'],
    ['E2E-08', 'Announcement flow', '1) Admin posts a needs-ack announcement. 2) Employee gets notification, opens app, acknowledges.', 'Notification once; ack count on web increments.', 'P2'],
    ['E2E-09', 'Language switch mid-flow', 'Switch to Marathi while a request is pending.', 'Status lines and tiles translate immediately without losing state.', 'P3'],
]))

# ------------------------------------------------------------------ render
doc.add_heading('4. Android app test cases', 1)
para('Preconditions for all Android cases: build installed, user logged in (unless the case says otherwise), network on, GPS on, camera / location permission granted.', italic=True)
count = 0
for title, rows in sections:
    lvl = 1 if title[0] != '4' else 2
    doc.add_heading(title, lvl)
    table(H, rows, W)
    count += len(rows)

doc.add_heading('9. Known limitations and open points to confirm with the product team', 1)
bullets([
    'Visit notes, photo and voice upload need connectivity (no offline queue for them). Confirm expected behaviour (E2E-06).',
    'The "Today\'s plan" card appears only when a plan exists for the logged-in user. Plans are created on the web screen Visit plans.',
    'Open orders / last visit / task due definitions: open orders = customer sales with sale_status pending or ordered, or payment_status pending, due or partial; last visit = last done plan for that customer before today; task due = 1 when the plan has a task text.',
    'Approval and rejection notifications arrive within 30 minutes (polling), not instantly, until FCM is added.',
    'Area tracking depends on the device keeping the foreground service alive; some manufacturers kill it (see AND-GEO-06).',
    'Header "Keep going, <name>!" can be truncated for very long names.',
])

doc.add_heading('10. Defect report template', 1)
table(['Field', 'What to write'], [
    ['ID / Title', 'Short, specific title (screen + what is wrong).'],
    ['Build / device', 'App version, Android version, device model, language, network.'],
    ['Account', 'User name and role used.'],
    ['Steps to reproduce', 'Numbered steps starting from a clean state.'],
    ['Expected vs actual', 'One line each.'],
    ['Evidence', 'Screenshot or screen recording, logcat excerpt (for crashes include the "FATAL EXCEPTION" block), API request / response.'],
    ['Severity / priority', 'Critical / High / Medium / Low and P1 / P2 / P3.'],
    ['Linked test case', 'Test case ID, e.g. AND-VIS-03.'],
], [5, 20])

doc.add_heading('11. Test execution summary sheet', 1)
table(['Section', 'Total', 'Passed', 'Failed', 'Blocked', 'Not run', 'Tester / date'],
      [[t.split(' ', 1)[0] + ' ' + t.split(' ', 1)[1][:40], len(r), '', '', '', '', ''] for t, r in sections] + [['Total', count, '', '', '', '', '']],
      [9, 2.5, 2.5, 2.5, 2.5, 2.5, 5])

out = r'C:\wamp64\www\captured_attendance\docs\ElintOmAT_QA_Test_Cases.docx'
doc.save(out)
print('saved', out, 'cases', count)
